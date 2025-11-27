import logging
import json
import re
import uuid
import tempfile
import os
from datetime import datetime
from typing import Dict, List, Optional
from pydantic import BaseModel, Field
from langchain_core.tools import tool
from docx import Document
import asyncio

from src.ai_core.services import search_service
from src.ai_core.models.chunks import SearchResult
from src.services.minio_service import minio_service
from src.config import AgentConfig

logger = logging.getLogger("uvicorn.error")


def _save_debug_mapping(
    document_type: str,
    placeholders: List[str],
    user_info: str,
    ai_raw_response: str,
    parsed_mapping: Dict,
    success: bool,
    error_msg: str = "",
    parsing_method: str = "structured_output"
) -> None:
    """Save debug information về mapping process"""
    
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        debug_filename = f"mapping_debug_{timestamp}_{document_type.replace(' ', '_')}.json"
        debug_path = os.path.join("debug_logs/document_filler", debug_filename)
        
        debug_data = {
            "timestamp": datetime.now().isoformat(),
            "document_type": document_type,
            "placeholders_detected": placeholders,
            "placeholders_count": len(placeholders),
            "user_info": user_info,
            "user_info_length": len(user_info),
            "ai_raw_response": ai_raw_response,
            "ai_response_length": len(ai_raw_response),
            "parsed_mapping": parsed_mapping,
            "mapped_fields_count": len(parsed_mapping.get("mapping", {})),
            "missing_fields": parsed_mapping.get("missing_fields", []),
            "missing_fields_count": len(parsed_mapping.get("missing_fields", [])),
            "confidence": parsed_mapping.get("confidence", 0.0),
            "parsing_method": parsing_method,
            "success": success,
            "error_message": error_msg,
            "model_used": "gemini-2.5-flash"
        }
        
        # Ensure directory exists
        os.makedirs(os.path.dirname(debug_path), exist_ok=True)
        
        # Save debug file
        with open(debug_path, 'w', encoding='utf-8') as f:
            json.dump(debug_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Debug mapping saved to: {debug_path}")
        
    except Exception as e:
        logger.error(f"Failed to save debug mapping: {e}")


class MappingResult(BaseModel):
    """Pydantic model cho kết quả mapping thông tin vào mẫu đơn"""
    mapping: str = Field(
        default="{}", 
        description="JSON string chứa mapping từ placeholder name (key) đến giá trị cụ thể (value). Ví dụ: '{\"ho_ten\": \"Nguyễn Văn A\", \"ngay_sinh\": \"01/01/1990\"}'"
    )
    missing_fields: List[str] = Field(
        default_factory=list, 
        description="List các placeholder không thể điền được do thiếu thông tin. Ví dụ: ['dia_chi', 'so_dien_thoai']"
    )
    confidence: float = Field(
        default=0.8, 
        ge=0.0, 
        le=1.0,
        description="Độ tin cậy của mapping từ 0.0 đến 1.0. Ví dụ: 0.85 nghĩa là 85% tin cậy"
    )
    notes: str = Field(
        default="", 
        description="Ghi chú về quá trình mapping hoặc lưu ý đặc biệt. Ví dụ: 'Đã mapping thành công 5/7 trường'"
    )
    
    class Config:
        schema_extra = {
            "example": {
                "mapping": "{\"ho_ten\": \"Nguyễn Văn A\", \"ngay_sinh\": \"01/01/1990\", \"gioi_tinh\": \"Nam\", \"quoc_tich\": \"Việt Nam\"}",
                "missing_fields": ["dia_chi", "so_dien_thoai"],
                "confidence": 0.85,
                "notes": "Đã mapping thành công 4/6 trường"
            }
        }

@tool
async def generate_filled_document(
    document_type: str = Field(
        description="Loại mẫu đơn cần điền, ví dụ: 'khai sinh', 'ly hôn', 'đăng ký cư trú'"
    ),
    user_provided_info: str = Field(
        description="Thông tin người dùng cung cấp dưới dạng text tự nhiên để điền vào mẫu đơn"
    )
) -> dict:
    """
    Sử dụng AI để phân tích thông tin người dùng và tạo file mẫu đơn đã được điền sẵn thông tin.
    AI sẽ thông minh map thông tin người dùng với các placeholder trong mẫu đơn.
    """
    
    try:
        logger.info(f"Generating filled document for type: {document_type}")
        
        # Step 1: Tìm kiếm mẫu đơn có placeholder
        results: list[SearchResult] = await search_service.search(
            app_id=ToolsConfig. document_template_app_id,
            query=f"mẫu đơn {document_type}",
            top_k=3,
            score_threshold=0.3,
            filters=[
                {
                    "bool": {
                        "should": [
                            {"match_phrase": {"metadata.title": f"mẫu đơn {document_type}"}},
                            {"match_phrase": {"metadata.title": document_type}},
                            {"wildcard": {"metadata.document_metadata.file_name": "*.docx"}},
                            {"wildcard": {"metadata.document_metadata.file_name": "*.doc"}}
                        ],
                        "minimum_should_match": 1
                    }
                }
            ]
        )
        
        if not results:
            return {
                "success": False,
                "message": f"Không tìm thấy mẫu đơn {document_type}",
                "suggestion": "Vui lòng thử với tên khác hoặc liên hệ hỗ trợ"
            }
        
        best_result = results[0]
        
        # Step 2: Lấy thông tin placeholder và file path
        if not (hasattr(best_result, 'document_metadata') and best_result.document_metadata):
            return {
                "success": False,
                "message": "Mẫu đơn không có thông tin metadata phù hợp"
            }
        
        metadata = best_result.document_metadata
        if not isinstance(metadata, dict):
            return {
                "success": False,
                "message": "Metadata mẫu đơn không đúng định dạng"
            }
        
        placeholder_file_path = metadata.get("file_path_placeholder", "")
        content_placeholder = metadata.get("content_placeholder", "")
        
        if not placeholder_file_path or not content_placeholder:
            return {
                "success": False,
                "message": "Mẫu đơn này chưa hỗ trợ tự động điền thông tin",
                "suggestion": "Vui lòng sử dụng mẫu đơn gốc và tự điền thông tin"
            }
        
        # Step 3: Extract placeholders từ content
        placeholders = re.findall(r'\{([^}]+)\}', content_placeholder)
        unique_placeholders = list(set(placeholders))
        
        logger.info(f"Found placeholders: {unique_placeholders}")
        
        # Step 4: Sử dụng AI để mapping thông tin
        mapping_result = await _ai_map_user_info_to_placeholders(
            user_provided_info, 
            unique_placeholders,
            document_type
        )
        
        if not mapping_result["success"]:
            return mapping_result
        
        # Step 5: Kiểm tra thông tin thiếu
        if mapping_result["missing_fields"]:
            missing_list = ", ".join(mapping_result["missing_fields"])
            return {
                "success": False,
                "message": f"Cần bổ sung thêm thông tin: {missing_list}",
                "missing_fields": mapping_result["missing_fields"],
                "current_mapping": mapping_result["mapping"],
                "suggestion": "Vui lòng cung cấp đầy đủ thông tin trên để tôi có thể tạo mẫu đơn hoàn chỉnh"
            }
        
        # Step 6: Download file placeholder và fill thông tin
        filled_file_result = await _download_and_fill_document(
            placeholder_file_path,
            mapping_result["mapping"],
            document_type,
            best_result.title
        )
        
        return filled_file_result
        
    except Exception as e:
        logger.error(f"Error in generate_filled_document: {e}")
        return {
            "success": False,
            "message": f"Đã có lỗi xảy ra: {str(e)}"
        }


async def _ai_map_user_info_to_placeholders(
    user_info: str, 
    placeholders: List[str], 
    document_type: str
) -> Dict:
    """Sử dụng Gemini AI để mapping thông tin người dùng với placeholders với enhanced error handling"""
    
    try:
        # Tạo prompt cải tiến với examples
        placeholders_text = "\n".join([f"- {p}" for p in placeholders])
        
        prompt = f"""
Bạn là AI chuyên gia điền mẫu đơn hành chính Việt Nam.

LOẠI MẪU ĐƠN: {document_type}

CÁC TRƯỜNG CẦN ĐIỀN:
{placeholders_text}

THÔNG TIN NGƯỜI DÙNG:
{user_info}

EXAMPLE OUTPUT:
{{"mapping": {{"ho_ten": "Nguyễn Văn A", "ngay_sinh": "12/03/1985", "dia_chi": "123 Lý Thường Kiệt, Hoàn Kiếm, Hà Nội"}}, "missing_fields": [], "confidence": 0.95, "notes": "Mapped thành công"}}

QUY TẮC:
- Phân tích thông minh và map chính xác
- Format ngày: dd/mm/yyyy
- Chỉ map khi chắc chắn 
- Liệt kê missing_fields nếu thiếu thông tin quan trọng

OUTPUT (chỉ JSON thuần túy):
"""
        
        logger.info(f"Calling AI mapping for {document_type} with {len(placeholders)} placeholders")
        logger.info(f"Prompt length: {len(prompt)}")
        logger.info(f"Prompt preview: {prompt[:300]}...")
        
        ai_raw_response = ""
        parsed_mapping = {}
        parsing_method = "structured_output"
        
        # Try structured output first
        try:
            gemini_model = AgentConfig.models["gemini-2.5-flash"]
            
            # Convert prompt to message format
            from langchain_core.messages import HumanMessage
            messages = [HumanMessage(content=prompt)]
            
            # Get raw response first để debug
            raw_response = await gemini_model.ainvoke(messages)
            ai_raw_response = raw_response.content if hasattr(raw_response, 'content') else str(raw_response)
            
            logger.info(f"Raw structured response length: {len(ai_raw_response)}")
            logger.info(f"Raw structured response preview: {ai_raw_response[:200]}...")
            
            # Try structured output
            structured_model = gemini_model.with_structured_output(MappingResult)
            
            if structured_model:
                logger.info("Using structured output approach")
                result = await structured_model.ainvoke(messages)
                
                logger.info(f"Structured result mapping: {result.mapping}")
                logger.info(f"Structured result confidence: {result.confidence}")
                
                # Parse JSON string từ result.mapping
                try:
                    mapping_dict = json.loads(result.mapping) if result.mapping else {}
                except json.JSONDecodeError:
                    logger.warning(f"Failed to parse mapping JSON string: {result.mapping}")
                    mapping_dict = {}
                
                # Check if result has meaningful data
                if mapping_dict or len(result.missing_fields) < len(placeholders):
                    parsed_mapping = {
                        "mapping": mapping_dict,
                        "missing_fields": result.missing_fields,
                        "confidence": result.confidence,
                        "notes": result.notes
                    }
                    
                    # Save debug info
                    _save_debug_mapping(
                        document_type, placeholders, user_info, 
                        ai_raw_response, parsed_mapping, 
                        True, "", "structured_output"
                    )
                    
                    return {
                        "success": True,
                        "mapping": mapping_dict,
                        "missing_fields": result.missing_fields,
                        "confidence": result.confidence,
                        "notes": result.notes
                    }
                else:
                    logger.warning("Structured output returned empty mapping, falling back to manual parsing")
                    parsing_method = "manual_parsing_after_empty_structured"
                    
        except Exception as structured_error:
            logger.warning(f"Structured parsing failed: {structured_error}")
            parsing_method = "manual_parsing"
        
        # Fallback to manual parsing
        try:
            gemini_model = AgentConfig.models["gemini-2.5-flash"]
            
            # Validate prompt before sending
            if not prompt or len(prompt.strip()) == 0:
                raise ValueError("Empty prompt provided for AI mapping")
            
            # Convert to message format for manual parsing too
            from langchain_core.messages import HumanMessage
            messages = [HumanMessage(content=prompt)]
            
            response = await gemini_model.ainvoke(messages)
            
            if hasattr(response, 'content'):
                ai_raw_response = response.content
            else:
                ai_raw_response = str(response)
            
            logger.info(f"Raw AI response length: {len(ai_raw_response)}")
            logger.info(f"Raw AI response preview: {ai_raw_response[:200]}...")
            
            # Enhanced JSON parsing with multiple strategies
            parsed_result = _parse_ai_response(ai_raw_response)
            
            if parsed_result:
                parsed_mapping = parsed_result
                
                # Save debug info
                _save_debug_mapping(
                    document_type, placeholders, user_info,
                    ai_raw_response, parsed_mapping, 
                    True, "", "manual_parsing"
                )
                
                return {
                    "success": True,
                    "mapping": parsed_result.get("mapping", {}),
                    "missing_fields": parsed_result.get("missing_fields", []),
                    "confidence": parsed_result.get("confidence", 0.0),
                    "notes": parsed_result.get("notes", "")
                }
            else:
                logger.warning("AI response parsing failed, using fallback extraction")
                parsing_method = "fallback_extraction"
                
        except Exception as ai_error:
            logger.error(f"AI call failed: {ai_error}")
            ai_raw_response = f"AI_ERROR: {str(ai_error)}"
            parsing_method = "fallback_extraction"
        
        # Fallback extraction khi AI thất bại
        fallback_result = _fallback_extraction(user_info, placeholders)
        parsed_mapping = fallback_result
        
        # Save debug info cho fallback
        _save_debug_mapping(
            document_type, placeholders, user_info,
            ai_raw_response, parsed_mapping, 
            False, "AI parsing failed, used fallback", parsing_method
        )
        
        return {
            "success": len(fallback_result.get("mapping", {})) > 0,
            "message": "AI mapping thất bại, sử dụng fallback extraction" if len(fallback_result.get("mapping", {})) == 0 else "Sử dụng fallback extraction",
            "mapping": fallback_result.get("mapping", {}),
            "missing_fields": fallback_result.get("missing_fields", placeholders),
            "confidence": 0.3,
            "notes": "Fallback extraction được sử dụng do AI parsing thất bại"
        }
            
    except Exception as e:
        logger.error(f"Critical error in AI mapping: {e}")
        return {
            "success": False,
            "message": f"Lỗi hệ thống khi xử lý AI: {str(e)}"
        }


def _parse_ai_response(response_content: str) -> Optional[Dict]:
    """Parse AI response với multiple strategies"""
    
    # Strategy 1: Direct JSON parse
    try:
        return json.loads(response_content.strip())
    except json.JSONDecodeError:
        pass
    
    # Strategy 2: Extract JSON block từ text
    try:
        # Tìm JSON block trong response
        json_pattern = r'\{.*?\}'
        json_match = re.search(json_pattern, response_content, re.DOTALL)
        if json_match:
            return json.loads(json_match.group())
    except json.JSONDecodeError:
        pass
    
    # Strategy 3: Extract từng trường một
    try:
        mapping = {}
        missing_fields = []
        
        # Extract mapping
        mapping_match = re.search(r'"mapping":\s*\{([^}]*)\}', response_content)
        if mapping_match:
            mapping_text = mapping_match.group(1)
            # Parse key-value pairs
            for match in re.finditer(r'"([^"]+)":\s*"([^"]+)"', mapping_text):
                mapping[match.group(1)] = match.group(2)
        
        # Extract missing_fields
        missing_match = re.search(r'"missing_fields":\s*\[([^\]]*)\]', response_content)
        if missing_match:
            missing_text = missing_match.group(1)
            missing_fields = [m.strip('"') for m in missing_text.split(',') if m.strip()]
        
        if mapping or missing_fields:
            return {
                "mapping": mapping,
                "missing_fields": missing_fields,
                "confidence": 0.7,
                "notes": "Extracted using regex parsing"
            }
    except Exception as e:
        logger.error(f"Regex parsing failed: {e}")
    
    return None


def _fallback_extraction(user_info: str, placeholders: List[str]) -> Dict:
    """Fallback extraction sử dụng regex patterns cơ bản"""
    
    mapping = {}
    
    # Common patterns cho thông tin cá nhân
    patterns = {
        "ho_ten": r"(?:tên|họ tên|tôi là|tôi tên)\s*:?\s*([A-ZÀÁẠẢÃÂẦẤẬẨẪĂẰẮẶẲẴÈÉẸẺẼÊỀẾỆỂỄÌÍỊỈĨÒÓỌỎÕÔỒỐỘỔỖƠỜỚỢỞỠÙÚỤỦŨƯỪỨỰỬỮỲÝỴỶỸĐ][a-zA-Zàáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ\s]+)",
        "ngay_sinh": r"(?:sinh|sinh ngày|ngày sinh)\s*:?\s*(\d{1,2}[\/\-]\d{1,2}[\/\-]\d{4})",
        "cmnd": r"(?:cmnd|cccd|số)\s*:?\s*(\d{9,12})",
        "dia_chi": r"(?:địa chỉ|ở|sống tại)\s*:?\s*([^,\n]+)"
    }
    
    user_info_lower = user_info.lower()
    
    for placeholder in placeholders:
        placeholder_lower = placeholder.lower()
        
        # Try exact match first
        if placeholder_lower in patterns:
            match = re.search(patterns[placeholder_lower], user_info, re.IGNORECASE)
            if match:
                mapping[placeholder] = match.group(1).strip()
        
        # Try fuzzy matching
        else:
            for pattern_name, pattern in patterns.items():
                if pattern_name in placeholder_lower or any(word in placeholder_lower for word in pattern_name.split('_')):
                    match = re.search(pattern, user_info, re.IGNORECASE)
                    if match:
                        mapping[placeholder] = match.group(1).strip()
                        break
    
    missing_fields = [p for p in placeholders if p not in mapping]
    
    return {
        "mapping": mapping,
        "missing_fields": missing_fields
    }


async def _download_and_fill_document(
    placeholder_file_path: str,
    mapping: Dict[str, str],
    document_type: str,
    document_title: str
) -> Dict:
    """Download file placeholder và fill thông tin"""
    
    try:
        # Auto-detect bucket
        bucket_name = "dataset"
        actual_file_path = placeholder_file_path
        
        if "/" in placeholder_file_path:
            potential_bucket = placeholder_file_path.split("/")[0]
            common_buckets = ["dataset", "documents", "default", "maudon"]
            if potential_bucket in common_buckets:
                bucket_name = potential_bucket
                actual_file_path = "/".join(placeholder_file_path.split("/")[1:])
        
        # Download file từ Minio - sử dụng signature đúng
        downloaded_path = await minio_service.download_file(
            actual_file_path,
            bucket_name
        )
        
        if not downloaded_path:
            return {
                "success": False,
                "message": "Không thể tải file mẫu đơn placeholder"
            }
        
        # Load và xử lý Word document
        doc = Document(downloaded_path)
        
        # Replace placeholders trong paragraphs
        for paragraph in doc.paragraphs:
            for placeholder, value in mapping.items():
                if f"{{{placeholder}}}" in paragraph.text:
                    paragraph.text = paragraph.text.replace(f"{{{placeholder}}}", str(value))
        
        # Replace placeholders trong tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for placeholder, value in mapping.items():
                        if f"{{{placeholder}}}" in cell.text:
                            cell.text = cell.text.replace(f"{{{placeholder}}}", str(value))
        
        # Tạo filename mới
        unique_id = str(uuid.uuid4())[:8]
        filled_filename = f"{document_type}_filled_{unique_id}.docx"
        filled_temp_path = os.path.join(tempfile.gettempdir(), filled_filename)
        
        # Save file đã điền
        doc.save(filled_temp_path)
        
        # Upload file mới lên Minio
        upload_path = f"filled_documents/{filled_filename}"
        upload_success = await minio_service.upload_file(
            filled_temp_path,
            upload_path,
            bucket_name
        )
        
        # Cleanup temp files
        try:
            os.unlink(downloaded_path)
            os.unlink(filled_temp_path)
        except:
            pass
        
        if not upload_success:
            return {
                "success": False,
                "message": "Không thể upload file đã điền lên hệ thống"
            }
        
        # Tạo download URL
        download_url = await minio_service.get_file_url(
            upload_path,
            bucket_name,
            expires_hours=24
        )
        
        if download_url:
            return {
                "success": True,
                "message": f"Mẫu đơn {document_type} đã được điền thông tin thành công",
                "document_title": document_title,
                "file_name": filled_filename,
                "download_url": download_url,
                "expires_in": "24 giờ",
                "type": "filled_document",
                "filled_fields": len(mapping),
                "instructions": [
                    "File đã được điền sẵn thông tin bạn cung cấp",
                    "Vui lòng kiểm tra và chỉnh sửa nếu cần",
                    "Mở file bằng Microsoft Word để xem chi tiết",
                    "Link download có hiệu lực trong 24 giờ"
                ]
            }
        else:
            return {
                "success": False,
                "message": "Không thể tạo link download cho file đã điền"
            }
            
    except Exception as e:
        logger.error(f"Error in download and fill document: {e}")
        return {
            "success": False,
            "message": f"Lỗi khi xử lý file: {str(e)}"
        }
