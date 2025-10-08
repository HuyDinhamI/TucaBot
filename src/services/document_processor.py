import os
import json
import re
import tempfile
import logging
from typing import Optional, Dict, Any
from docx import Document

logger = logging.getLogger("uvicorn.error")

class DocumentProcessor:
    """Service xử lý và tạo placeholder cho mẫu đơn docx"""
    
    def __init__(self, llm_client=None):
        self.llm_client = llm_client
    
    def extract_text_blocks(self, doc_path: str) -> list[str]:
        """Trích xuất các đoạn văn bản không rỗng từ file Word"""
        try:
            doc = Document(doc_path)
            blocks = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    blocks.append(text)
            return blocks
        except Exception as e:
            logger.error(f"Error extracting text from {doc_path}: {e}")
            return []
    
    def clean_and_parse_json(self, raw_text: str) -> dict:
        """Làm sạch output AI và parse JSON an toàn"""
        try:
            cleaned = raw_text.strip()
            
            # Bỏ ```json / ```
            cleaned = re.sub(r"^```[a-zA-Z]*", "", cleaned)
            cleaned = re.sub(r"```$", "", cleaned).strip()
            
            # Thay "{abc}" -> "{{abc}}"
            cleaned = re.sub(r'"\{(.*?)\}"', r'"{{\1}}"', cleaned)
            
            return json.loads(cleaned)
        except Exception as e:
            logger.error(f"JSON parse failed: {e}")
            logger.error(f"Raw cleaned text: {cleaned}")
            return {}
    
    async def generate_placeholders_with_ai(self, text_blocks: list[str]) -> dict:
        """Sử dụng AI để tạo mapping placeholder cho các text blocks"""
        if not self.llm_client or not text_blocks:
            return {}
        
        prompt = f"""
Bạn là hệ thống xử lý biểu mẫu. 
Hãy tìm các chỗ người dùng cần điền thông tin trong các đoạn sau 
và sinh ra placeholder theo cú pháp {{snake_case}}.

Yêu cầu:
- Không tạo placeholder cho phần tiêu đề (CỘNG HÒA, TỜ KHAI...) hoặc "Chú thích", "Ví dụ".
- Nếu có nhiều field trên cùng 1 dòng (ví dụ Giới tính / Dân tộc / Quốc tịch), hãy tạo placeholder riêng cho từng field.
- Nếu có dấu .... hoặc …… thì bỏ chúng đi, thay hẳn bằng placeholder.
- Chỉ trả về JSON hợp lệ, không giải thích, không có ```json.

Đầu ra JSON dạng: {{
   "original_text": "new_text_with_placeholder"
}}

Các đoạn cần phân tích:
{text_blocks}
"""

        try:
            response = await self.llm_client.chat.completions.create(
                model="gpt-4o-mini",  # Sử dụng model phù hợp
                messages=[
                    {"role": "system", "content": "Bạn là trợ lý AI chuyên tạo placeholder từ biểu mẫu."},
                    {"role": "user", "content": prompt}
                ],
                temperature=0,
            )
            
            raw_response = response.choices[0].message.content.strip()
            return self.clean_and_parse_json(raw_response)
            
        except Exception as e:
            logger.error(f"Error generating placeholders with AI: {e}")
            return {}
    
    def create_simple_placeholders(self, text_blocks: list[str]) -> dict:
        """Tạo placeholder đơn giản không cần AI (fallback method)"""
        mapping = {}
        
        for text in text_blocks:
            # Tìm các dấu chấm liên tiếp (ít nhất 3 dấu)
            if re.search(r"\.{3,}", text):
                # Xác định loại placeholder dựa trên context
                lower_text = text.lower()
                
                if any(keyword in lower_text for keyword in ['họ tên', 'tên', 'name']):
                    new_text = re.sub(r"\.{3,}", "{{ho_ten}}", text)
                elif any(keyword in lower_text for keyword in ['ngày', 'date', 'sinh']):
                    new_text = re.sub(r"\.{3,}", "{{ngay_thang_nam}}", text)
                elif any(keyword in lower_text for keyword in ['địa chỉ', 'address', 'nơi ở']):
                    new_text = re.sub(r"\.{3,}", "{{dia_chi}}", text)
                elif any(keyword in lower_text for keyword in ['số điện thoại', 'phone', 'dt']):
                    new_text = re.sub(r"\.{3,}", "{{so_dien_thoai}}", text)
                elif any(keyword in lower_text for keyword in ['cmnd', 'cccd', 'số']):
                    new_text = re.sub(r"\.{3,}", "{{so_cmnd}}", text)
                else:
                    new_text = re.sub(r"\.{3,}", "{{thong_tin}}", text)
                
                mapping[text] = new_text
        
        return mapping
    
    def replace_placeholders_in_document(self, doc_path: str, mapping: dict, output_path: str) -> bool:
        """Thay thế các field trong Word document theo mapping và lưu file mới"""
        try:
            doc = Document(doc_path)
            
            for paragraph in doc.paragraphs:
                for original_text, new_text in mapping.items():
                    if original_text in paragraph.text:
                        paragraph.text = paragraph.text.replace(original_text, new_text)
            
            # Lưu file mới
            doc.save(output_path)
            logger.info(f"Document saved with placeholders: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"Error replacing placeholders in document: {e}")
            return False
    
    async def process_document(
        self, 
        doc_path: str, 
        use_ai: bool = True
    ) -> Optional[str]:
        """
        Xử lý mẫu đơn và tạo placeholder
        
        Args:
            doc_path: Đường dẫn file docx gốc
            use_ai: Có sử dụng AI để tạo placeholder không
            
        Returns:
            Đường dẫn file đã xử lý hoặc None nếu thất bại
        """
        try:
            # Trích xuất text blocks
            text_blocks = self.extract_text_blocks(doc_path)
            if not text_blocks:
                logger.warning(f"No text blocks found in {doc_path}")
                return None
            
            # Tạo mapping placeholder
            if use_ai and self.llm_client:
                mapping = await self.generate_placeholders_with_ai(text_blocks)
            else:
                mapping = self.create_simple_placeholders(text_blocks)
            
            if not mapping:
                logger.warning("No placeholders generated")
                # Trả về file gốc nếu không có placeholder nào
                return doc_path
            
            # Tạo file output
            temp_dir = tempfile.mkdtemp()
            filename = os.path.basename(doc_path)
            name, ext = os.path.splitext(filename)
            output_filename = f"{name}_placeholder{ext}"
            output_path = os.path.join(temp_dir, output_filename)
            
            # Thay thế placeholder và lưu file
            success = self.replace_placeholders_in_document(doc_path, mapping, output_path)
            
            if success:
                logger.info(f"Document processed successfully: {output_path}")
                return output_path
            else:
                return None
                
        except Exception as e:
            logger.error(f"Error processing document {doc_path}: {e}")
            return None

# Singleton instance
document_processor = DocumentProcessor()
